#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
SUSPICIOUS LINK STATUS CHECK - to run on an ISOLATED sandbox / VM.

Reads the LINKS sheet, checks columns C (reported site) and D (redirect),
fills column G (ΤΡΕΧΟΥΣΑ ΚΑΤΑΣΤΑΣΗ) with ACTIVE / OFFLINE and writes a full
log (HTTP status, redirect chain, final destination, IP, title).

How to use it:
    python check_links.py <input.xlsx> [output.xlsx]

Warning:
  - Do NOT run this from the Service network without an outbound proxy/VPN or from a physical machine.
  - With --head-only the page body is not downloaded (less noisy).
  - Tracking links (mlsend, sendgrid, powr.io, bit.ly) are burned on first
    open; use --no-redirect if you only want the first hop.
"""

# import libraries
import sys
import csv
import socket
import argparse
import datetime
import urllib.parse

import requests
import urllib3
import openpyxl

# disable potential worning from printing
urllib3.disable_warnings()

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")


# fixed variables that can be tuned
TIMEOUT = 15
COL_URL = 3        # C
COL_REDIRECT = 4   # D
COL_STATUS = 7     # G



def normalize(u):
    if not u:
        return None
    u = str(u).strip()
    if not u or "@" in u.split("/")[0] or u.lower() in ("not found", "-"):
        return None            # email address ή placeholder, όχι URL
    if not u.startswith(("http://", "https://")):
        u = "http://" + u
    return u

def resolve_ip(url):
    try:
        host = urllib.parse.urlparse(url).hostname
        return socket.gethostbyname(host) if host else ""
    except Exception:
        return ""

def probe(url, args):
    """Επιστρέφει dict με το αποτέλεσμα του ελέγχου."""
    res = {"url": url, "status": "", "http": "", "final": "", "chain": "",
           "ip": "", "title": "", "error": ""}
    res["ip"] = resolve_ip(url)
    if not res["ip"]:
        res["status"] = "OFFLINE"
        res["error"] = "DNS resolution failed / NXDOMAIN"
        return res

    method = "HEAD" if args.head_only else "GET"
    try:
        r = requests.request(
            method, url,
            timeout=TIMEOUT,
            allow_redirects=not args.no_redirect,
            verify=False,
            headers={"User-Agent": UA, "Accept-Language": "el-GR,el;q=0.9,en;q=0.8"},
        )
        res["http"] = r.status_code
        res["final"] = r.url
        res["chain"] = " -> ".join(h.url for h in r.history)
        if not args.head_only and r.text:
            body = r.text
            i = body.lower().find("<title")
            if i != -1:
                j = body.find(">", i)
                k = body.lower().find("</title>", j)
                if j != -1 and k != -1:
                    res["title"] = body[j + 1:k].strip()[:120]
        res["status"] = "ACTIVE" if r.status_code < 400 else "OFFLINE"
    except requests.exceptions.SSLError as e:
        res["status"] = "ACTIVE"          # απαντά ο host, πρόβλημα μόνο στο TLS
        res["error"] = "SSL: %s" % str(e)[:150]
    except requests.exceptions.ConnectionError as e:
        res["status"] = "OFFLINE"
        res["error"] = "CONN: %s" % str(e)[:150]
    except requests.exceptions.Timeout:
        res["status"] = "OFFLINE"
        res["error"] = "TIMEOUT"
    except Exception as e:
        res["status"] = "OFFLINE"
        res["error"] = str(e)[:150]
    return res

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output", nargs="?")
    ap.add_argument("--sheet", default="LINKS")
    ap.add_argument("--head-only", action="store_true")
    ap.add_argument("--no-redirect", action="store_true")
    ap.add_argument("--log", default="link_check_log.csv")
    args = ap.parse_args()

    out = args.output or args.input.replace(".xlsx", "_CHECKED.xlsx")
    wb = openpyxl.load_workbook(args.input)
    ws = wb[args.sheet]

    logf = open(args.log, "w", newline="", encoding="utf-8-sig")
    log = csv.writer(logf, delimiter=";")
    log.writerow(["row", "field", "url", "status", "http", "ip",
                  "redirect_chain", "final_url", "title", "error", "checked_at"])

    for row in range(2, ws.max_row + 1):
        verdicts = []
        for col, label in ((COL_URL, "C"), (COL_REDIRECT, "D")):
            url = normalize(ws.cell(row=row, column=col).value)
            if not url:
                continue
            r = probe(url, args)
            verdicts.append(r["status"])
            log.writerow([row, label, r["url"], r["status"], r["http"], r["ip"],
                          r["chain"], r["final"], r["title"], r["error"],
                          datetime.datetime.now().isoformat(timespec="seconds")])
            print("%-4s %-8s %-6s %s" % (row, r["status"], r["http"], r["url"][:90]))
        if verdicts:
            ws.cell(row=row, column=COL_STATUS,
                    value="ACTIVE" if "ACTIVE" in verdicts else "OFFLINE")

    logf.close()
    wb.save(out)
    print("\nOK ->", out, "|", args.log)


if __name__ == "__main__":
    main()
"""
URL Checker — XLSX upload, έλεγχος συνδέσμων, εξαγωγή Excel.

    pip install fastapi uvicorn python-multipart pandas openpyxl httpx
    export URLSCAN_API_KEY=...
    uvicorn main:app --reload --port 8001

UI:  http://localhost:8001/
"""

import datetime as dt
import io
import math
import pathlib
import re
import tempfile
from typing import Optional

import numpy as np
import pandas as pd
from fastapi import FastAPI, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from checker import check_urls
from redirect_check import check_redirects

STATIC = pathlib.Path(__file__).parent / "static"

app = FastAPI(title="URL Checker")
app.mount("/static", StaticFiles(directory=STATIC), name="static")

ALLOWED = (".xlsx", ".xlsm", ".xls")
MAX_BYTES = 25 * 1024 * 1024

URL_COLS = [
    "URL ή ΙΡ ΚΑΤΑΓΓΕΛΛΟΜΕΝΟΥ ΙΣΤΟΤΟΠΟY",
    "URL ή ΙΡ ΑΝΑΚΑΤΕΥΘΥΝΣΗΣ",
]
PRIMARY_COL = URL_COLS[0]
STATUS_COL = "ΤΡΕΧΟΥΣΑ ΚΑΤΑΣΤΑΣΗ (ACTIVE/OFFLINE/BROWSER BLOCKED)"

# STATE του checker -> λεκτικό που χρησιμοποιείται ήδη στα αρχεία
STATUS_MAP = {
    "ACTIVE":          "Active",
    "REDIRECT":        "Active",
    "NOT FOUND (404)": "Page Not Found Er 404",
    "GONE":            "Offline",
    "FORBIDDEN":       "Forbidden (403)",
    "SERVER ERROR":    "Server Error",
    "DNS/NXDOMAIN":    "Offline",
    "OFFLINE":         "Offline",
    "TIMEOUT":         "Offline",
    "SSL ERROR":       "Active (SSL error)",
    "EMPTY":           None,
}

EXPORT_EXTRA = [
    f"{PRIMARY_COL} :: CODE",
    "REDIRECTED",
    "REDIRECT CHAIN",
    "REDIRECT HOPS",
    f"{URL_COLS[1]} :: URL",
    f"{URL_COLS[1]} :: STATE",
]

FILL = {
    "Active":                PatternFill("solid", fgColor="E3F2E8"),
    "Offline":               PatternFill("solid", fgColor="F8E6E4"),
    "Page Not Found Er 404": PatternFill("solid", fgColor="F8E6E4"),
}

_REFANG = [
    (re.compile(r"h(?:xx|tt)p(s?)\s*\[?\s*:\s*\]?\s*/\s*/\s*\]?", re.I), r"http\1://"),
    (re.compile(r"\[\.\]|\(\.\)|\[dot\]", re.I), "."),
    (re.compile(r"\[@\]|\(at\)|\[at\]", re.I), "@"),
]


def refang(v):
    if not isinstance(v, str):
        return v
    for pat, rep in _REFANG:
        v = pat.sub(rep, v)
    return v.strip()


def jsonify(v):
    """Μετατροπή κελιού σε τιμή που δέχεται το json.dumps."""
    if v is None or v is pd.NaT:
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, (np.floating, float)):
        f = float(v)
        return f if math.isfinite(f) else None
    if isinstance(v, np.bool_):
        return bool(v)
    if isinstance(v, (pd.Timestamp, dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, pd.Timedelta):
        return str(v)
    if isinstance(v, bytes):
        return v.decode("utf-8", "replace")
    if isinstance(v, (str, int, bool)):
        return v
    return str(v)


def _cells(records, cols):
    """(δείκτης γραμμής, στήλη, refanged url) για κάθε μη κενό κελί URL."""
    return [
        (i, c, refang(rec[c]))
        for i, rec in enumerate(records)
        for c in cols
        if isinstance(rec.get(c), str) and rec[c].strip()
    ]


def apply_status(records, df):
    """Γεμίζει τη στήλη ΤΡΕΧΟΥΣΑ ΚΑΤΑΣΤΑΣΗ από το αποτέλεσμα του ελέγχου."""
    if STATUS_COL not in [str(c) for c in df.columns]:
        return
    state_key = f"{PRIMARY_COL} :: STATE"
    for rec in records:
        st = rec.get(state_key)
        if st is None:
            continue
        # κράτα ό,τι έχει συμπληρωθεί χειροκίνητα:
        # if rec.get(STATUS_COL): continue
        rec[STATUS_COL] = STATUS_MAP.get(st, st)


def ordered_columns(records, original):
    """Αρχικές στήλες στη σειρά τους, μετά οι στήλες ελέγχου."""
    extra = [k for k in dict.fromkeys(k for r in records for k in r)
             if k not in original]
    return original + extra


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/upload")
async def upload(
    file: UploadFile = File(...),
    sheet: Optional[str] = Query(None, description="Όνομα φύλλου· κενό = όλα"),
    max_rows: int = Query(100, ge=1, le=10000),
    check: bool = Query(False, description="Άμεσος έλεγχος (γρήγορος, από τη δική σας IP)"),
    scan: bool = Query(False, description="urlscan.io (αργό, βρίσκει JS ανακατευθύνσεις)"),
    visibility: str = Query("private", pattern="^(private|unlisted|public)$"),
):
    if not file.filename.lower().endswith(ALLOWED):
        raise HTTPException(400, f"Μη αποδεκτός τύπος αρχείου: {file.filename}")

    raw = await file.read()
    if len(raw) > MAX_BYTES:
        raise HTTPException(413, "Το αρχείο είναι πολύ μεγάλο")

    try:
        sheets = pd.read_excel(
            io.BytesIO(raw),
            sheet_name=sheet if sheet else None,
            engine="openpyxl",
        )
    except Exception as e:
        raise HTTPException(422, f"Αδυναμία ανάγνωσης: {e}")

    if isinstance(sheets, pd.DataFrame):
        sheets = {sheet: sheets}

    payload = {}
    for name, df in sheets.items():
        print(f"\n=== {file.filename} :: φύλλο '{name}' "
              f"({df.shape[0]} γραμμές x {df.shape[1]} στήλες) ===")

        records = [
            {str(k): jsonify(v) for k, v in row.items()}
            for row in df.head(max_rows).to_dict(orient="records")
        ]

        # --- άμεσος έλεγχος ---------------------------------------------
        if check:
            targets = _cells(records, [str(c) for c in URL_COLS
                                       if str(c) in [str(x) for x in df.columns]])
            results = await check_urls([u for _, _, u in targets])
            for (i, c, _), res in zip(targets, results):
                records[i][f"{c} :: URL"] = res["url"]
                records[i][f"{c} :: STATE"] = res["state"]
                records[i][f"{c} :: CODE"] = res["status"]
                records[i][f"{c} :: FINAL"] = res["final_url"]
                print(f"  [check] {res['state']:<16} {res['status']}  {res['url'][:90]}")

        # --- urlscan.io --------------------------------------------------
                if scan:
                    if PRIMARY_COL not in [str(c) for c in df.columns]:
                        raise HTTPException(422, f"Δεν βρέθηκε η στήλη: {PRIMARY_COL}")
                    targets = _cells(records, [PRIMARY_COL])
                    results = await check_redirects([u for _, _, u in targets])
                    for (i, _, _), res in zip(targets, results):
                        records[i]["REDIRECTED"] = res["redirected"]
                        records[i]["FINAL URL"] = res["final_url"]
                        records[i]["REDIRECT CHAIN"] = res["redirect_chain"]
                        records[i]["REDIRECT HOPS"] = res["redirect_hops"]
                        records[i]["FINAL STATUS"] = res["status_code"]
                        records[i]["SERVER"] = res["server"]
                        records[i]["LOOP"] = res["loop"]
                        records[i]["REDIRECT ERROR"] = res["error"]
                        flag = "REDIR" if res["redirected"] else "  -  "
                        print(f"  [redir] {flag} {res['redirect_hops']} hops  "
                            f"{(res['submitted_url'] or '')[:70]}")

        # --- συμπλήρωση ΤΡΕΧΟΥΣΑ ΚΑΤΑΣΤΑΣΗ -------------------------------
        apply_status(records, df)

        original = [str(c) for c in df.columns]
        payload[str(name)] = {
            "rows": int(df.shape[0]),
            "cols": int(df.shape[1]),
            "columns": original,
            "enriched_columns": [c for c in ordered_columns(records, original)
                                 if c not in original],
            "data": records,
        }

    return {"filename": file.filename, "sheets": payload}


@app.post("/upload/xlsx")
async def upload_xlsx(
    file: UploadFile = File(...),
    sheet: Optional[str] = Query(None),
    max_rows: int = Query(10000, ge=1, le=10000),
    check: bool = Query(True),
    scan: bool = Query(False),
    visibility: str = Query("private", pattern="^(private|unlisted|public)$"),
):
    """Ίδιος έλεγχος, αλλά επιστρέφει έτοιμο .xlsx με τις αρχικές στήλες."""
    result = await upload(file=file, sheet=sheet, max_rows=max_rows,
                          check=check, scan=scan, visibility=visibility)

    tmp = pathlib.Path(tempfile.mkdtemp()) / "excel_προς_ΕΑΚ.xlsx"

    with pd.ExcelWriter(tmp, engine="openpyxl") as w:
        for name, s in result["sheets"].items():
            cols = s["columns"] + [c for c in EXPORT_EXTRA
                                   if c in s["enriched_columns"]]
            pd.DataFrame(s["data"], columns=cols).to_excel(
                w, sheet_name=str(name)[:31], index=False)

    wb = load_workbook(tmp)
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]]
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, size=10)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 45
        ws.freeze_panes = "A2"
        for i, h in enumerate(headers, 1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = 55 if "URL" in str(h) else 26
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Font(name="Arial", size=10)
                c.alignment = Alignment(vertical="top")
        if STATUS_COL in headers:
            col = headers.index(STATUS_COL) + 1
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value in FILL:
                    cell.fill = FILL[cell.value]
    wb.save(tmp)

    return FileResponse(
        tmp,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=tmp.name,
    )
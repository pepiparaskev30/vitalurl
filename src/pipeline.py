import os
import re
import unicodedata
import warnings
from functools import partial
from concurrent.futures import ThreadPoolExecutor

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from utilities import read_json_dictionary, check_row, to_greek, unwrap_if_safelink, unwrap_safelink, refang

warnings.filterwarnings("ignore", message="Unverified HTTPS request")


# κωδικας ροης 
OUTPUT_COLUMNS = [
    "ΗΜΕΡΟΜΗΝΙΑ",
    "URL ή ΙΡ ΚΑΤΑΓΓΕΛΛΟΜΕΝΟΥ ΙΣΤΟΤΟΠΟY",
    "URL ή ΙΡ ΑΝΑΚΑΤΕΥΘΥΝΣΗΣ",
    "ΤΡΟΠΟΣ ΕΝΗΜΕΡΩΣΗΣ",
    "ΕΙΔΟΣ ΑΠΕΙΛΗΣ /ΗΛΕΚΤΡΟΝΙΚΗΣ ΑΠΑΤΗΣ - ΠΑΡΑΠΛΑΝΗΣΗΣ",
    "ΤΡΕΧΟΥΣΑ ΚΑΤΑΣΤΑΣΗ (ACTIVE/OFFLINE/BROWSER BLOCKED)",
    "ΜΙΜΟΥΜΕΝΟΣ ΙΣΤΟΤΟΠΟΣ",
]
DATE_COL = OUTPUT_COLUMNS[0]
URL_COL = OUTPUT_COLUMNS[1]
REDIRECT_COL = OUTPUT_COLUMNS[2]
STATUS_COL = OUTPUT_COLUMNS[5]

# active = ο ιστότοπος λειτουργεί, down = εκτός λειτουργίας, error = αβέβαιο αποτέλεσμα
DOWN_PREFIXES = ("SUSPENDED", "NOT FOUND", "GONE", "DNS")

BRAND_COL = "ΜΙΜΟΥΜΕΝΟΣ ΙΣΤΟΤΟΠΟΣ"
OTHER_GROUP = "Λοιπά"
EMPTY_GROUP = "Χωρίς τιμή"


def load_config(src_path):
    config = {
        "ua_agents": read_json_dictionary(src_path, "agents.json"),
        "dictionary_columns": read_json_dictionary(src_path, "dictionary.json"),
        "status_gr": read_json_dictionary(src_path, "status_gr.json"),
        "brand_categories": compile_brand_rules(
            read_json_dictionary(src_path, "brand_categories.json") or {}
        ),
    }
    if not config["ua_agents"] or not config["dictionary_columns"]:
        raise RuntimeError(f"Could not load agents.json or dictionary.json from {src_path}")
    return config


def normalize(text):
    """Πεζά, χωρίς τόνους και τελείες, ς → σ — για ανεκτική σύγκριση."""
    text = unicodedata.normalize("NFD", str(text).lower())
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return text.replace("ς", "σ").replace(".", "").strip()


def compile_brand_rules(rules):
    """{"Τράπεζες": ["eurobank", ...]} → [("Τράπεζες", regex), ...]"""
    compiled = []
    for group, keywords in rules.items():
        words = sorted({normalize(k) for k in keywords if str(k).strip()}, key=len, reverse=True)
        if words:
            pattern = r"(?<!\w)(?:" + "|".join(re.escape(w) for w in words) + r")(?!\w)"
            compiled.append((group, re.compile(pattern)))
    return compiled


def brand_group(brand, url, rules):
    """Κατηγορία φορέα από τον μιμούμενο ιστότοπο· αν είναι κενός, από το URL."""
    for text in (brand, url):
        if text is None or pd.isna(text) or not str(text).strip():
            continue
        norm = normalize(text)
        for group, pattern in rules:
            if pattern.search(norm):
                return group
        if text is brand:
            return OTHER_GROUP
    return EMPTY_GROUP


def defang_column(df, col):
    """Defanged μορφή (hxxps[://]...) ώστε οι σύνδεσμοι να μην είναι clickable στο Excel."""
    df[col] = df[col].apply(lambda x: defang(str(x)) if pd.notna(x) and str(x).strip() else x)


def category(status):
    if status is None:
        return "empty"
    if status.startswith("ACTIVE"):
        return "active"
    if status.startswith(DOWN_PREFIXES):
        return "down"
    return "error"


def process_df(old_df, ua_agents, dictionary_columns, status_gr=None, workers=20,
               check=True, redirects=True, max_rows=None, brand_categories=None):
    """Returns (new_df, categories, groups) — one status category and one brand group per row.

    check=False     → only rename columns and unwrap Safe Links, no requests are sent
    redirects=False → the redirect column keeps its original value
    """
    if max_rows:
        old_df = old_df.head(max_rows)

    new_df = old_df.rename(columns=dictionary_columns).reindex(columns=OUTPUT_COLUMNS)
    new_df[URL_COL] = new_df[URL_COL].apply(unwrap_if_safelink)

    # τυχόν ανακατευθύνσεις από το αρχικό αρχείο → κανονική μορφή (refang + unwrap Safe Links)
    new_df[REDIRECT_COL] = new_df[REDIRECT_COL].apply(
        lambda x: unwrap_safelink(refang(x)) if pd.notna(x) and str(x).strip() else x
    )

    rules = brand_categories or []
    groups = [brand_group(b, u, rules) for b, u in zip(new_df[BRAND_COL], new_df[URL_COL])]

    if not check:
        defang_column(new_df, REDIRECT_COL)
        return new_df, ["empty"] * len(new_df), groups

    rows = new_df.to_dict("records")
    with ThreadPoolExecutor(max_workers=workers) as ex:
        results = list(ex.map(partial(check_row, user_agents=ua_agents), rows))

    categories = [category(status) for status, _ in results]
    new_df[STATUS_COL] = [to_greek(status, status_gr) for status, _ in results]
    if redirects:
        new_df[REDIRECT_COL] = [
            redirect if redirect else existing
            for (_, redirect), existing in zip(results, new_df[REDIRECT_COL])
        ]
    defang_column(new_df, REDIRECT_COL)
    return new_df, categories, groups


STATUS_STYLE = {
    "active": ("E3F2E8", "1F6F43"),
    "down":   ("F8E6E4", "8A2D2D"),
    "error":  ("FAF0D9", "8A6116"),
}


def save_excel(df, categories, path):
    df.to_excel(path, index=False)
    wb = load_workbook(path)
    ws = wb.active

    widths = {DATE_COL: 14, URL_COL: 60, REDIRECT_COL: 60, STATUS_COL: 34}
    for i, col in enumerate(df.columns, 1):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = widths.get(col, 26)
        h = ws.cell(1, i)
        h.font = Font(name="Arial", bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor="16191D")
        h.alignment = Alignment(wrap_text=True, vertical="center")

    status_idx = list(df.columns).index(STATUS_COL) + 1
    date_idx = list(df.columns).index(DATE_COL) + 1
    for r, cat in enumerate(categories, start=2):
        for c in range(1, len(df.columns) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name="Arial")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, date_idx).number_format = "DD/MM/YYYY"
        if cat in STATUS_STYLE:
            fill, color = STATUS_STYLE[cat]
            cell = ws.cell(r, status_idx)
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(name="Arial", bold=True, color=color)

    ws.freeze_panes = "A2"
    wb.save(path)
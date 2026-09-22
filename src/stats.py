import json
from collections import Counter, defaultdict

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from pipeline import normalize, BRAND_COL, OTHER_GROUP, EMPTY_GROUP, STATUS_STYLE

TYPE_COL = "ΕΙΔΟΣ ΑΠΕΙΛΗΣ /ΗΛΕΚΤΡΟΝΙΚΗΣ ΑΠΑΤΗΣ - ΠΑΡΑΠΛΑΝΗΣΗΣ"

CAT_ORDER = ["active", "down", "error", "empty"]
CAT_LABELS = {
    "active": "Ενεργά",
    "down": "Εκτός λειτουργίας",
    "error": "Αβέβαιο",
    "empty": "Χωρίς έλεγχο",
}
TOTAL = "Σύνολο"
PCT = "% Ενεργά"
GROUP = "Κατηγορία"
BRAND = "Μιμούμενος ιστότοπος"
TYPE = "Είδος απειλής"


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def canonical_labels(values):
    """Ενοποιεί γραφές που διαφέρουν μόνο σε πεζά/τόνους/τελείες.

    'Eurobank', 'EUROBANK', 'eurobank ' → ένα κλειδί, με ετικέτα τη συχνότερη γραφή.
    """
    seen = defaultdict(Counter)
    keys = []
    for v in values:
        if v is None or pd.isna(v) or not str(v).strip():
            keys.append(None)
            continue
        text = str(v).strip()
        key = normalize(text)
        seen[key][text] += 1
        keys.append(key)
    label = {k: c.most_common(1)[0][0] for k, c in seen.items()}
    return [label[k] if k else EMPTY_GROUP for k in keys]


def _table(index_cols, categories):
    """Πίνακας συχνοτήτων: γραμμές = index_cols, στήλες = κατηγορίες κατάστασης."""
    df = pd.DataFrame(index_cols)
    df["_cat"] = categories
    keys = list(index_cols.keys())

    t = (df.groupby(keys + ["_cat"]).size()
           .unstack("_cat", fill_value=0)
           .reindex(columns=CAT_ORDER, fill_value=0))
    t.insert(0, TOTAL, t[CAT_ORDER].sum(axis=1))
    checked = t[TOTAL] - t["empty"]
    t[PCT] = (t["active"] / checked * 100).round(1).where(checked > 0)
    t = t.rename(columns=CAT_LABELS).reset_index()
    t.columns.name = None
    return t


def _total_row(t, label_col):
    tot = {c: "" for c in t.columns}
    tot[label_col] = TOTAL
    for c in [TOTAL] + list(CAT_LABELS.values()):
        tot[c] = int(t[c].sum())
    checked = tot[TOTAL] - tot[CAT_LABELS["empty"]]
    tot[PCT] = round(tot[CAT_LABELS["active"]] / checked * 100, 1) if checked else None
    return pd.concat([t, pd.DataFrame([tot])], ignore_index=True)


def _order_groups(t, group_order):
    rank = {g: i for i, g in enumerate(group_order)}
    return t.assign(_o=t[GROUP].map(lambda g: rank.get(g, len(rank))))


# ---------------------------------------------------------------------------
# Public
# ---------------------------------------------------------------------------

def build_stats(df, categories, groups, group_order):
    """Returns {"by_group", "by_brand", "by_type"} as DataFrames.

    group_order: η σειρά των κατηγοριών όπως στο brand_categories.json
    """
    order = list(group_order) + [OTHER_GROUP, EMPTY_GROUP]
    brands = canonical_labels(df[BRAND_COL])
    types = canonical_labels(df[TYPE_COL])

    # 1. Ανά κατηγορία φορέα
    by_group = _table({GROUP: groups}, categories)
    by_group = (_order_groups(by_group, order)
                  .sort_values(["_o"]).drop(columns="_o").reset_index(drop=True))
    by_group = _total_row(by_group, GROUP)

    # 2. Ανά κατηγορία → μιμούμενο ιστότοπο
    by_brand = _table({GROUP: groups, BRAND: brands}, categories)
    by_brand = (_order_groups(by_brand, order)
                  .sort_values(["_o", TOTAL], ascending=[True, False])
                  .drop(columns="_o").reset_index(drop=True))

    # 3. Ανά κατηγορία → είδος απειλής
    by_type = _table({GROUP: groups, TYPE: types}, categories)
    by_type = (_order_groups(by_type, order)
                 .sort_values(["_o", TOTAL], ascending=[True, False])
                 .drop(columns="_o").reset_index(drop=True))

    return {"by_group": by_group, "by_brand": by_brand, "by_type": by_type}


def stats_to_json(stats):
    return {k: json.loads(v.to_json(orient="records", force_ascii=False)) for k, v in stats.items()}


SHEETS = [
    ("by_group", "Ανά κατηγορία"),
    ("by_brand", "Ανά ιστότοπο"),
    ("by_type", "Ανά είδος απειλής"),
]

HEADER_FILL = PatternFill("solid", fgColor="16191D")
TOTAL_FILL = PatternFill("solid", fgColor="ECEAE5")
THIN = Side(style="thin", color="DCD9D2")


def save_stats_excel(stats, path, subtitle=""):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for key, sheet in SHEETS:
            stats[key].to_excel(xw, sheet_name=sheet, index=False, startrow=2)

    wb = load_workbook(path)
    status_cols = {CAT_LABELS[k]: k for k in ("active", "down", "error")}

    for key, sheet in SHEETS:
        ws = wb[sheet]
        t = stats[key]
        ncols = len(t.columns)

        ws["A1"] = f"Στατιστικά {sheet.lower()}"
        ws["A1"].font = Font(name="Arial", bold=True, size=13)
        ws["A2"] = subtitle
        ws["A2"].font = Font(name="Arial", color="6B737D")

        for i, col in enumerate(t.columns, 1):
            c = ws.cell(3, i)
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
            c.alignment = Alignment(wrap_text=True, vertical="center",
                                    horizontal="left" if col in (GROUP, BRAND, TYPE) else "center")
            width = 30 if col in (GROUP, BRAND, TYPE) else 13
            ws.column_dimensions[c.column_letter].width = width

        for r in range(4, 4 + len(t)):
            is_total = ws.cell(r, 1).value == TOTAL
            for i, col in enumerate(t.columns, 1):
                c = ws.cell(r, i)
                c.font = Font(name="Arial", bold=is_total)
                c.border = Border(bottom=THIN)
                if col not in (GROUP, BRAND, TYPE):
                    c.alignment = Alignment(horizontal="center")
                if col == PCT:
                    c.number_format = '0.0"%"'
                if is_total:
                    c.fill = TOTAL_FILL
                elif col in status_cols and c.value:
                    fill, color = STATUS_STYLE[status_cols[col]]
                    c.fill = PatternFill("solid", fgColor=fill)
                    c.font = Font(name="Arial", color=color, bold=True)

        ws.row_dimensions[3].height = 30
        ws.freeze_panes = "A4"

    # Γράφημα: στοιβαγμένες μπάρες ανά κατηγορία (χωρίς τη γραμμή συνόλου)
    ws = wb["Ανά κατηγορία"]
    t = stats["by_group"]
    n = len(t) - 1
    if n > 0:
        cols = list(t.columns)
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.title = "Κατάσταση ανά κατηγορία φορέα"
        chart.y_axis.title = "Πλήθος URL"
        chart.height = 7 + n * 0.6
        chart.width = 18
        for key, color in (("active", "1F6F43"), ("down", "8A2D2D"), ("error", "C8962E")):
            idx = cols.index(CAT_LABELS[key]) + 1
            data = Reference(ws, min_col=idx, min_row=3, max_row=3 + n)
            chart.add_data(data, titles_from_data=True)
            chart.series[-1].graphicalProperties.solidFill = color
            chart.series[-1].graphicalProperties.line.solidFill = color
        chart.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n))
        chart.x_axis.scaling.orientation = "maxMin"   # ίδια σειρά με τον πίνακα
        chart.legend.position = "b"
        ws.add_chart(chart, f"A{6 + n}")

    wb.save(path)